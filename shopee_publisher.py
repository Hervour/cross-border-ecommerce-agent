"""Module 3: publish Shopee products through API v2 or mass-upload Excel.

Mode A uses the official Shopee Open API v2 signed request flow and creates
an unpublished product with ``item_status=UNLIST``. Mode B exports an Excel
workbook based on the common Shopee mass-upload columns for manual completion
in Seller Centre.
"""

from __future__ import annotations

import hashlib
import hmac
import json
import logging
import os
import re
import sys
import time
from pathlib import Path
from typing import Any, Mapping

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from ecommerce_agent.models import ProductVariation, StandardProductListing

LOGGER = logging.getLogger(__name__)

ADD_ITEM_PATH = "/api/v2/product/add_item"
DEFAULT_BASE_URL = "https://partner.shopeemobile.com"
DRAFT_ITEM_STATUS = "UNLIST"


class ShopeePublisherError(Exception):
    """Base exception for Shopee publishing failures."""


class ShopeeConfigurationError(ShopeePublisherError):
    """Raised when required Shopee credentials are missing or invalid."""


class ShopeeValidationError(ShopeePublisherError):
    """Raised when a listing cannot be converted into a valid API payload."""


class ShopeeAPIError(ShopeePublisherError):
    """A structured Shopee API error with a user-friendly explanation."""

    def __init__(
        self,
        message: str,
        *,
        error_code: str | None = None,
        request_id: str | None = None,
        status_code: int | None = None,
        response_data: Mapping[str, Any] | None = None,
    ) -> None:
        super().__init__(message)
        self.error_code = error_code
        self.request_id = request_id
        self.status_code = status_code
        self.response_data = dict(response_data or {})


def generate_shopee_signature(
    partner_id: int,
    partner_key: str,
    path: str,
    timestamp: int,
    access_token: str | None = None,
    shop_id: int | None = None,
) -> str:
    """Generate the HMAC-SHA256 signature for a Shopee Open API v2 request.

    Public API base string: ``partner_id + path + timestamp``.
    Shop API base string adds ``access_token + shop_id``.
    """

    if not isinstance(partner_id, int) or partner_id <= 0:
        raise ShopeeConfigurationError("partner_id must be a positive integer")
    if not partner_key:
        raise ShopeeConfigurationError("partner_key cannot be empty")
    if not path.startswith("/"):
        raise ShopeeValidationError("Shopee API path must start with '/'")
    if not isinstance(timestamp, int) or timestamp <= 0:
        raise ShopeeValidationError("timestamp must be a positive Unix timestamp")

    base_string = f"{partner_id}{path}{timestamp}"
    if access_token is not None or shop_id is not None:
        if not access_token or not isinstance(shop_id, int) or shop_id <= 0:
            raise ShopeeConfigurationError(
                "access_token and shop_id must both be valid for a Shop API request"
            )
        base_string += f"{access_token}{shop_id}"

    return hmac.new(
        partner_key.encode("utf-8"),
        base_string.encode("utf-8"),
        hashlib.sha256,
    ).hexdigest()

class ShopeeAPIClient:
    """Minimal Shopee Open API v2 client for product creation."""

    def __init__(
        self,
        *,
        partner_id: int,
        partner_key: str,
        access_token: str,
        shop_id: int,
        base_url: str = DEFAULT_BASE_URL,
        session: requests.Session | None = None,
        timeout: float = 30.0,
    ) -> None:
        if not partner_key:
            raise ShopeeConfigurationError("partner_key cannot be empty")
        if not access_token:
            raise ShopeeConfigurationError("access_token cannot be empty")
        if not isinstance(partner_id, int) or partner_id <= 0:
            raise ShopeeConfigurationError("partner_id must be a positive integer")
        if not isinstance(shop_id, int) or shop_id <= 0:
            raise ShopeeConfigurationError("shop_id must be a positive integer")

        self.partner_id = partner_id
        self.partner_key = partner_key
        self.access_token = access_token
        self.shop_id = shop_id
        self.base_url = base_url.rstrip("/")
        self.timeout = timeout
        self.session = session or requests.Session()

    @classmethod
    def from_env(cls) -> "ShopeeAPIClient":
        """Build a client from SHOPEE_* environment variables."""

        required = {
            "SHOPEE_PARTNER_ID": os.environ.get("SHOPEE_PARTNER_ID"),
            "SHOPEE_PARTNER_KEY": os.environ.get("SHOPEE_PARTNER_KEY"),
            "SHOPEE_ACCESS_TOKEN": os.environ.get("SHOPEE_ACCESS_TOKEN"),
            "SHOPEE_SHOP_ID": os.environ.get("SHOPEE_SHOP_ID"),
        }
        missing = [key for key, value in required.items() if not value]
        if missing:
            raise ShopeeConfigurationError(
                "Missing Shopee API environment variables: " + ", ".join(missing)
            )
        try:
            partner_id = int(str(required["SHOPEE_PARTNER_ID"]))
            shop_id = int(str(required["SHOPEE_SHOP_ID"]))
        except ValueError as exc:
            raise ShopeeConfigurationError(
                "SHOPEE_PARTNER_ID and SHOPEE_SHOP_ID must be integers"
            ) from exc

        return cls(
            partner_id=partner_id,
            partner_key=str(required["SHOPEE_PARTNER_KEY"]),
            access_token=str(required["SHOPEE_ACCESS_TOKEN"]),
            shop_id=shop_id,
            base_url=os.environ.get("SHOPEE_BASE_URL", DEFAULT_BASE_URL),
            timeout=float(os.environ.get("SHOPEE_TIMEOUT", "30")),
        )

    def generate_signature(self, path: str, timestamp: int | None = None) -> str:
        """Generate a Shop API signature for the configured shop credentials."""

        return generate_shopee_signature(
            partner_id=self.partner_id,
            partner_key=self.partner_key,
            path=path,
            timestamp=timestamp or int(time.time()),
            access_token=self.access_token,
            shop_id=self.shop_id,
        )

    def upload_as_draft(self, listing_data: dict[str, Any]) -> dict[str, Any]:
        """Convert a standard listing and submit it as an unpublished item."""

        payload = build_shopee_item_payload(listing_data)
        LOGGER.info(
            "Submitting Shopee draft item '%s' with category_id=%s",
            payload.get("item_name"),
            payload.get("category_id"),
        )
        try:
            return self.add_item(payload)
        except ShopeeAPIError as exc:
            friendly = friendly_shopee_error(exc.error_code, str(exc))
            LOGGER.error("Shopee add_item failed: %s", friendly)
            raise ShopeeAPIError(
                friendly,
                error_code=exc.error_code,
                request_id=exc.request_id,
                status_code=exc.status_code,
                response_data=exc.response_data,
            ) from exc
        except requests.RequestException as exc:
            friendly = "Shopee interface connection failed; check network, endpoint, and timeout settings."
            LOGGER.error("%s Details: %s", friendly, exc)
            raise ShopeeAPIError(friendly) from exc

    def add_item(self, payload: dict[str, Any]) -> dict[str, Any]:
        """Call Shopee Open API v2 ``v2.product.add_item``."""

        return self._request(ADD_ITEM_PATH, payload)

    def call_api(self, path: str, payload: dict[str, Any]) -> dict[str, Any]:
        """Call another signed Shopee Open API v2 endpoint."""

        return self._request(path, payload)

    def _request(self, path: str, payload: dict[str, Any]) -> dict[str, Any]:
        timestamp = int(time.time())
        params = {
            "partner_id": self.partner_id,
            "timestamp": timestamp,
            "access_token": self.access_token,
            "shop_id": self.shop_id,
            "sign": self.generate_signature(path, timestamp),
        }
        url = f"{self.base_url}{path}"

        try:
            response = self.session.request(
                "POST",
                url,
                params=params,
                json=payload,
                timeout=self.timeout,
                headers={"Content-Type": "application/json"},
            )
            response.raise_for_status()
        except requests.RequestException as exc:
            raise ShopeeAPIError(
                f"Shopee request failed: {exc}",
                status_code=getattr(exc.response, "status_code", None),
            ) from exc

        try:
            data = response.json()
        except ValueError as exc:
            raise ShopeeAPIError(
                "Shopee returned a non-JSON response",
                status_code=response.status_code,
            ) from exc

        if not isinstance(data, dict):
            raise ShopeeAPIError(
                "Shopee returned an unexpected response shape",
                status_code=response.status_code,
            )

        error_code = data.get("error")
        if error_code not in (None, "", 0, "0"):
            raise ShopeeAPIError(
                str(data.get("error_msg") or "Shopee API request failed"),
                error_code=str(error_code),
                request_id=data.get("request_id"),
                status_code=response.status_code,
                response_data=data,
            )
        return data


def friendly_shopee_error(error_code: str | None, message: str) -> str:
    """Translate common Shopee validation errors into actionable guidance."""

    combined = f"{error_code or ''} {message}".lower()
    if "categor" in combined or "attribute" in combined:
        return (
            "Shopee category or required attributes are invalid/missing. "
            "Call the category and attribute APIs, then fill category_id and attribute_list."
        )
    if "image" in combined or "photo" in combined or "dimension" in combined:
        return (
            "Shopee image validation failed. Upload compliant images through the media API "
            "and pass their image IDs in image.image_id_list."
        )
    if "name" in combined or "title" in combined or "length" in combined:
        return (
            "Shopee rejected the product title. Keep it within the category limit, remove "
            "duplicate keywords, and avoid prohibited characters."
        )
    if "price" in combined or "currency" in combined:
        return (
            "Shopee rejected the price. Check the shop currency, category price range, and "
            "the minimum/maximum allowed price."
        )
    if "logistic" in combined or "shipping" in combined:
        return (
            "Shopee logistics configuration is missing or disabled. Provide a valid "
            "logistic_info list for the target shop."
        )
    if "token" in combined or "sign" in combined or "partner" in combined:
        return (
            "Shopee authentication failed. Refresh access_token and verify partner_id, "
            "partner_key, shop_id, timestamp, and signature generation."
        )
    return f"Shopee API error: {message}"

def _plain_text(value: str) -> str:
    text = re.sub(r"^#{1,6}\s*", "", value, flags=re.MULTILINE)
    text = re.sub(r"[*_`]+", "", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _coerce_standard_listing(listing_data: dict[str, Any]) -> StandardProductListing:
    if not isinstance(listing_data, dict):
        raise ShopeeValidationError("listing_data must be a dictionary")

    standard_fields = (
        "title",
        "bullet_points",
        "description",
        "category",
        "tags",
        "variations",
    )
    try:
        return StandardProductListing.model_validate(
            {key: listing_data[key] for key in standard_fields if key in listing_data}
        )
    except (ValidationError, KeyError) as exc:
        raise ShopeeValidationError(f"Invalid StandardProductListing data: {exc}") from exc


def _normalize_image_ids(listing_data: dict[str, Any]) -> list[str]:
    raw_value: Any = listing_data.get("image_id_list")
    if raw_value is None:
        raw_value = listing_data.get("image_list")

    image = listing_data.get("image")
    if raw_value is None and isinstance(image, dict):
        raw_value = image.get("image_id_list")
    if raw_value is None and isinstance(image, list):
        raw_value = image
    if isinstance(raw_value, str):
        raw_value = [raw_value]
    if not isinstance(raw_value, list):
        return []

    image_ids: list[str] = []
    for item in raw_value:
        if isinstance(item, dict):
            item = item.get("image_id") or item.get("id") or item.get("url")
        if item in (None, ""):
            continue
        value = str(item).strip()
        if value.startswith(("http://", "https://")):
            raise ShopeeValidationError(
                "image_list must contain Shopee image IDs, not public image URLs. "
                "Upload images through Shopee media API first."
            )
        image_ids.append(value)
    return list(dict.fromkeys(image_ids))


def _variation_stock(
    listing_data: dict[str, Any], variation: ProductVariation, index: int, default: int
) -> int:
    per_variation = listing_data.get("variation_stocks")
    if isinstance(per_variation, dict):
        value = per_variation.get(variation.en_spec) or per_variation.get(str(index))
        if value is not None:
            return max(1, int(value))
    if isinstance(per_variation, list) and index < len(per_variation):
        return max(1, int(per_variation[index]))
    return default


def build_shopee_item_payload(listing_data: dict[str, Any]) -> dict[str, Any]:
    """Convert module-2 data into the Shopee v2.product.add_item request body."""

    listing = _coerce_standard_listing(listing_data)

    category_id = listing_data.get("category_id")
    if category_id in (None, ""):
        raise ShopeeValidationError(
            "category_id is required for Shopee API upload. The category text produced "
            "by module 2 cannot be used as an API category ID."
        )
    try:
        category_id = int(category_id)
    except (TypeError, ValueError) as exc:
        raise ShopeeValidationError("category_id must be an integer") from exc

    image_ids = _normalize_image_ids(listing_data)
    if not image_ids:
        raise ShopeeValidationError(
            "image_id_list is required for Shopee API upload. Upload compliant images "
            "through Shopee media API before calling add_item."
        )

    default_stock = max(1, int(listing_data.get("normal_stock", listing_data.get("stock", 10))))
    min_price = min(variation.target_price_sgd for variation in listing.variations)
    bullet_text = "\n".join(f"• {_plain_text(item)}" for item in listing.bullet_points)
    description = f"{bullet_text}\n\n{_plain_text(listing.description)}".strip()

    item_status = str(listing_data.get("item_status") or DRAFT_ITEM_STATUS).upper()
    if item_status not in {"UNLIST", "NORMAL"}:
        raise ShopeeValidationError("item_status must be UNLIST or NORMAL")
    if item_status == "NORMAL":
        LOGGER.warning("item_status=NORMAL will publish the item; the default draft mode is UNLIST")

    weight = float(listing_data.get("weight", listing_data.get("package_weight", 0.1)))
    logistic_info = listing_data.get("logistic_info") or listing_data.get("logistics") or []
    attribute_list = listing_data.get("attribute_list") or []

    payload: dict[str, Any] = {
        "category_id": category_id,
        "item_name": listing.title,
        "description": description,
        "original_price": round(min_price, 2),
        "normal_stock": default_stock,
        "weight": weight,
        "item_status": item_status,
        "logistic_info": logistic_info,
        "image": {"image_id_list": image_ids},
        "brand": listing_data.get(
            "brand", {"brand_id": 0, "original_brand_name": "No Brand"}
        ),
        "item_dangerous": int(listing_data.get("item_dangerous", 0)),
    }

    if attribute_list:
        payload["attribute_list"] = attribute_list

    if len(listing.variations) > 1:
        payload["tier_variation"] = [
            {
                "name": str(listing_data.get("variation_name") or "Style"),
                "option_list": [
                    {"option": variation.en_spec} for variation in listing.variations
                ],
            }
        ]
        payload["model_list"] = [
            {
                "tier_index": [index],
                "original_price": round(variation.target_price_sgd, 2),
                "normal_stock": _variation_stock(
                    listing_data, variation, index, default_stock
                ),
            }
            for index, variation in enumerate(listing.variations)
        ]
    elif listing_data.get("sku"):
        payload["item_sku"] = str(listing_data["sku"])

    return payload


def upload_as_draft(
    listing_data: dict[str, Any],
    *,
    client: ShopeeAPIClient | None = None,
) -> dict[str, Any]:
    """Upload one listing as an unpublished Shopee item."""

    api_client = client or ShopeeAPIClient.from_env()
    return api_client.upload_as_draft(listing_data)

MASS_UPLOAD_HEADERS = [
    "Category ID",
    "Category",
    "Product Name",
    "Product Description",
    "SKU",
    "Variation Name",
    "Variation Option",
    "Price",
    "Stock",
    "Weight (kg)",
    "Cover Image",
    "Image 1",
    "Image 2",
    "Image 3",
    "Image 4",
    "Image 5",
    "Tags",
]


def _style_mass_upload_sheet(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="EE4D2D")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "A": 14,
        "B": 36,
        "C": 42,
        "D": 60,
        "E": 24,
        "F": 18,
        "G": 30,
        "H": 12,
        "I": 10,
        "J": 12,
        "K": 42,
        "L": 42,
        "M": 42,
        "N": 42,
        "O": 42,
        "P": 42,
        "Q": 30,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24


def export_shopee_mass_upload_excel(
    listing: StandardProductListing, output_path: str
) -> str:
    """Export a Shopee mass-upload workbook with one row per variation.

    Shopee's exact template differs by market and category. This workbook uses
    the common upload columns and includes an Instructions sheet; export the
    official category template when Seller Centre provides one.
    """

    if not isinstance(listing, StandardProductListing):
        try:
            listing = StandardProductListing.model_validate(listing)
        except (ValidationError, TypeError) as exc:
            raise ShopeeValidationError(
                f"listing must be a StandardProductListing: {exc}"
            ) from exc

    target = Path(output_path)
    if target.suffix.lower() != ".xlsx":
        target = target.with_suffix(".xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)

    image_urls = list(getattr(listing, "image_urls", []) or [])
    image_columns = image_urls[:6] + [""] * max(0, 6 - len(image_urls))
    variation_name = "Style"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mass Upload"
    sheet.append(MASS_UPLOAD_HEADERS)

    for variation in listing.variations:
        row = [
            "",
            listing.category,
            listing.title,
            _plain_text(listing.description),
            variation.en_spec,
            variation_name,
            variation.en_spec,
            round(variation.target_price_sgd, 2),
            "",
            "",
            *image_columns,
            ", ".join(listing.tags),
        ]
        sheet.append(row)

    _style_mass_upload_sheet(sheet)

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Field", "Action required"])
    instructions.append([
        "Category ID",
        "Fill the numeric category ID from Shopee's category API or Seller Centre.",
    ])
    instructions.append([
        "Images",
        "Upload images first, or paste the URLs required by the exact Shopee template for your market.",
    ])
    instructions.append([
        "Weight (kg)",
        "Fill the packed parcel weight accepted by the selected logistics channel.",
    ])
    instructions.append([
        "Template compatibility",
        "Shopee mass-upload columns vary by market and category. Replace this workbook with the official downloaded template when available.",
    ])
    instructions.append([
        "API mode",
        "For API upload, image_id_list must contain Shopee media IDs, not public image URLs.",
    ])
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 100
    for cell in instructions[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="EE4D2D")
    for row in instructions.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    try:
        workbook.save(target)
    except OSError as exc:
        raise ShopeePublisherError(f"Could not write Excel file: {exc}") from exc

    LOGGER.info("Shopee mass-upload workbook written to %s", target.resolve())
    return str(target.resolve())


def _demo_listing() -> StandardProductListing:
    return StandardProductListing(
        title=(
            "Portable Folding Camping Chair Lightweight Aluminum Moon Chair "
            "for Outdoor Travel Beach Picnic"
        ),
        bullet_points=[
            "🪑 Compact folding frame fits easily in a car trunk",
            "🪶 Lightweight aluminum alloy frame for easy carrying",
            "🧵 Durable 600D Oxford fabric seat is easy to clean",
            "🎒 Side mesh pocket keeps a bottle and phone within reach",
            "🏕️ Suitable for camping, fishing, picnics, and outdoor concerts",
        ],
        description=(
            "## Portable Outdoor Camping Chair\n\n"
            "A lightweight folding chair for convenient outdoor seating.\n\n"
            "### Specifications\n"
            "- Frame: Aluminum alloy\n"
            "- Fabric: 600D Oxford\n"
            "- Open size: 52 x 46 x 68 cm"
        ),
        category="Sports & Outdoors > Camping & Hiking > Camping Furniture > Chairs",
        tags=["camping chair", "folding chair", "outdoor chair", "moon chair", "beach chair"],
        variations=[
            ProductVariation(
                original_spec="米白色 标准款",
                en_spec="Beige Standard",
                cost_cny=58.0,
                target_price_sgd=23.50,
            ),
            ProductVariation(
                original_spec="军绿色 加宽款",
                en_spec="Army Green Wide",
                cost_cny=69.0,
                target_price_sgd=27.30,
            ),
        ],
    )


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    demo_listing = _demo_listing()
    exported_path = export_shopee_mass_upload_excel(
        demo_listing,
        str(Path(__file__).with_name("shopee_mass_upload_demo.xlsx")),
    )
    print(f"Excel exported: {exported_path}")

    required_api_env = (
        "SHOPEE_PARTNER_ID",
        "SHOPEE_PARTNER_KEY",
        "SHOPEE_ACCESS_TOKEN",
        "SHOPEE_SHOP_ID",
        "SHOPEE_DEMO_CATEGORY_ID",
        "SHOPEE_DEMO_IMAGE_ID",
    )
    missing = [key for key in required_api_env if not os.environ.get(key)]
    if missing:
        print("Shopee API demo skipped. Missing: " + ", ".join(missing))
    else:
        api_listing_data = demo_listing.model_dump()
        api_listing_data.update(
            {
                "category_id": int(os.environ["SHOPEE_DEMO_CATEGORY_ID"]),
                "image_id_list": [os.environ["SHOPEE_DEMO_IMAGE_ID"]],
                "weight": 1.2,
                "logistic_info": [],
                "attribute_list": [],
            }
        )
        result = upload_as_draft(api_listing_data)
        print(json.dumps(result, indent=2, ensure_ascii=False, default=str))