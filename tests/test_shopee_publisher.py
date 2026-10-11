from __future__ import annotations

import hashlib
import hmac
import tempfile
import unittest
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

from ecommerce_agent.models import ProductVariation, StandardProductListing
from shopee_publisher import (
    ADD_ITEM_PATH,
    ShopeeAPIError,
    ShopeeAPIClient,
    ShopeeValidationError,
    build_shopee_item_payload,
    export_shopee_mass_upload_excel,
    generate_shopee_signature,
)


class _FakeResponse:
    def __init__(self, data: dict[str, Any], status_code: int = 200) -> None:
        self.data = data
        self.status_code = status_code

    def json(self) -> dict[str, Any]:
        return self.data

    def raise_for_status(self) -> None:
        if self.status_code >= 400:
            error = requests.HTTPError(f"HTTP {self.status_code}")
            error.response = self  # type: ignore[assignment]
            raise error


class _RecordingSession:
    def __init__(self, response_data: dict[str, Any]) -> None:
        self.response_data = response_data
        self.calls: list[dict[str, Any]] = []

    def request(self, method: str, url: str, **kwargs: Any) -> _FakeResponse:
        self.calls.append({"method": method, "url": url, "kwargs": kwargs})
        return _FakeResponse(self.response_data)


def _sample_listing() -> StandardProductListing:
    return StandardProductListing(
        title=(
            "Portable Folding Camping Chair Lightweight Aluminum Moon Chair "
            "for Outdoor Travel Beach Picnic"
        ),
        bullet_points=[
            "🪑 Compact folding frame fits easily in a car trunk",
            "🪶 Lightweight aluminum alloy frame for easy carrying",
            "🧵 Durable 600D Oxford fabric seat is easy to clean",
            "🎒 Side mesh pocket keeps a bottle and phone within reach",
            "🏕️ Suitable for camping, fishing, picnics, and outdoor concerts",
        ],
        description="## Portable Chair\nA lightweight folding chair for outdoor seating.",
        category="Sports & Outdoors > Camping Chairs",
        tags=["camping chair", "folding chair", "outdoor chair", "moon chair", "beach chair"],
        variations=[
            ProductVariation(
                original_spec="米白色",
                en_spec="Beige",
                cost_cny=58.0,
                target_price_sgd=23.50,
            ),
            ProductVariation(
                original_spec="军绿色",
                en_spec="Army Green",
                cost_cny=69.0,
                target_price_sgd=27.30,
            ),
        ],
    )


def _api_data() -> dict[str, Any]:
    data = _sample_listing().model_dump()
    data.update(
        {
            "category_id": 100017,
            "image_id_list": ["image-id-1"],
            "weight": 1.2,
            "normal_stock": 20,
            "logistic_info": [{"logistic_id": 1, "enabled": True}],
            "attribute_list": [{"attribute_id": 10, "attribute_value_list": [{"value": "Aluminum"}]}],
        }
    )
    return data


class SignatureTests(unittest.TestCase):
    def test_signature_matches_shopee_base_string(self) -> None:
        signature = generate_shopee_signature(
            partner_id=123,
            partner_key="secret",
            path=ADD_ITEM_PATH,
            timestamp=1_700_000_000,
            access_token="token",
            shop_id=456,
        )
        base_string = f"123{ADD_ITEM_PATH}1700000000token456"
        expected = hmac.new(
            b"secret", base_string.encode("utf-8"), hashlib.sha256
        ).hexdigest()
        self.assertEqual(signature, expected)


class PayloadTests(unittest.TestCase):
    def test_payload_maps_to_unpublished_add_item_request(self) -> None:
        payload = build_shopee_item_payload(_api_data())

        self.assertEqual(payload["category_id"], 100017)
        self.assertEqual(payload["item_status"], "UNLIST")
        self.assertEqual(payload["original_price"], 23.50)
        self.assertEqual(payload["normal_stock"], 20)
        self.assertEqual(payload["image"]["image_id_list"], ["image-id-1"])
        self.assertEqual(len(payload["model_list"]), 2)
        self.assertEqual(payload["tier_variation"][0]["name"], "Style")

    def test_public_image_url_is_rejected_for_api_mode(self) -> None:
        data = _api_data()
        data["image_id_list"] = ["https://example.com/product.jpg"]
        with self.assertRaises(ShopeeValidationError):
            build_shopee_item_payload(data)


class ShopeeClientTests(unittest.TestCase):
    def test_upload_as_draft_sends_signed_shop_api_request(self) -> None:
        session = _RecordingSession({"item_id": 9988})
        client = ShopeeAPIClient(
            partner_id=123,
            partner_key="secret",
            access_token="token",
            shop_id=456,
            base_url="https://partner.shopeemobile.com",
            session=session,  # type: ignore[arg-type]
        )

        result = client.upload_as_draft(_api_data())

        self.assertEqual(result["item_id"], 9988)
        call = session.calls[0]
        self.assertEqual(call["method"], "POST")
        self.assertTrue(call["url"].endswith(ADD_ITEM_PATH))
        self.assertEqual(call["kwargs"]["params"]["partner_id"], 123)
        self.assertEqual(call["kwargs"]["params"]["shop_id"], 456)
        self.assertIn("sign", call["kwargs"]["params"])
        self.assertEqual(call["kwargs"]["json"]["item_status"], "UNLIST")

    def test_api_validation_error_is_translated(self) -> None:
        session = _RecordingSession(
            {
                "error": "error_param",
                "error_msg": "image dimension is invalid",
                "request_id": "request-1",
            }
        )
        client = ShopeeAPIClient(
            partner_id=123,
            partner_key="secret",
            access_token="token",
            shop_id=456,
            session=session,  # type: ignore[arg-type]
        )

        with self.assertRaises(ShopeeAPIError) as context:
            client.upload_as_draft(_api_data())

        self.assertIn("image validation failed", str(context.exception).lower())
        self.assertEqual(context.exception.request_id, "request-1")


class ExcelExportTests(unittest.TestCase):
    def test_export_creates_mass_upload_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "shopee_upload.xlsx"
            exported = export_shopee_mass_upload_excel(
                _sample_listing(), str(output_path)
            )

            self.assertTrue(Path(exported).exists())
            workbook = load_workbook(exported)
            self.assertEqual(workbook.sheetnames, ["Mass Upload", "Instructions"])
            sheet = workbook["Mass Upload"]
            self.assertEqual(sheet["A1"].value, "Category ID")
            self.assertEqual(sheet["C1"].value, "Product Name")
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(sheet["G2"].value, "Beige")
            self.assertEqual(sheet["H2"].value, 23.5)


if __name__ == "__main__":
    unittest.main()